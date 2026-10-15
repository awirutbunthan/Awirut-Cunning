import re
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import ZipFile

import pdfplumber
from openpyxl import load_workbook

from app.domain.models import DocumentKind, LineItem

FIELD_PATTERNS = {
    "legal_name": r"(?:customer|buyer|company|ลูกค้า)\s*[:#-]?\s*(.+)",
    "taxpayer_id": r"(?:tax(?:payer)?\s*(?:id|no)?|เลขประจำตัวผู้เสียภาษี)\s*[:#-]?\s*([\d-]{8,20})",
    "document_number": r"(?:po|invoice|receipt)\s*(?:no|number|#)?\s*[:#-]?\s*([A-Za-z0-9/-]+)",
    "date": r"(?:date|วันที่)\s*[:#-]?\s*([0-9]{1,4}[/-][0-9]{1,2}[/-][0-9]{1,4})",
    "grand_total": r"(?:grand\s*total|total|รวมทั้งสิ้น)\s*[:#-]?\s*(?:THB|฿)?\s*([\d,]+(?:\.\d{1,2})?)",
}


def _decimal(value: str) -> Decimal | None:
    try:
        return Decimal(value.replace(",", ""))
    except (InvalidOperation, AttributeError):
        return None


def parse_text(
    text: str, source: str
) -> tuple[dict[str, str | None], list[LineItem], list[dict[str, int | str]]]:
    fields: dict[str, str | None] = {}
    provenance: list[dict[str, int | str]] = []
    for field, pattern in FIELD_PATTERNS.items():
        match = re.search(pattern, text, re.IGNORECASE | re.MULTILINE)
        fields[field] = match.group(1).strip() if match else None
        if match:
            provenance.append({"field": field, "source": source, "offset": match.start(1)})
    lines: list[LineItem] = []
    for row, line in enumerate(text.splitlines(), 1):
        match = re.match(r"\s*(.+?)\s*\|\s*(\d+(?:\.\d+)?)\s*\|\s*([\d,]+(?:\.\d+)?)\s*$", line)
        if match:
            quantity, price = _decimal(match.group(2)), _decimal(match.group(3))
            lines.append(
                LineItem(
                    match.group(1).strip(),
                    quantity,
                    price,
                    quantity * price if quantity and price else None,
                    {"source": source, "row": row},
                )
            )
    return fields, lines, provenance


def extract_pdf(
    data: bytes, filename: str
) -> tuple[str, dict[str, str | None], list[LineItem], list[dict[str, int | str]], str]:
    with pdfplumber.open(BytesIO(data)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    fields, lines, provenance = parse_text(text, filename)
    return text, fields, lines, provenance, "OCR_REQUIRED" if not text.strip() else "PROCESSED"


def extract_xlsx(
    data: bytes, filename: str
) -> tuple[str, dict[str, str | None], list[LineItem], list[dict[str, int | str]], str]:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    rows: list[str] = []
    provenance: list[dict[str, int | str]] = []
    for sheet in workbook.worksheets:
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
            values = [str(cell) for cell in row if cell is not None]
            if values:
                rows.append(" | ".join(values))
                provenance.append({"source": filename, "sheet": sheet.title, "row": row_index})
    text = "\n".join(rows)
    fields, lines, field_provenance = parse_text(text, filename)
    return text, fields, lines, provenance + field_provenance, "PROCESSED"


def extract_zip(
    data: bytes, filename: str
) -> tuple[str, dict[str, str | None], list[LineItem], list[dict[str, int | str]], str]:
    texts: list[str] = []
    provenance: list[dict[str, int | str]] = []
    with ZipFile(BytesIO(data)) as archive:
        for info in archive.infolist():
            if info.is_dir():
                continue
            name = PurePosixPath(info.filename).name
            extension = PurePosixPath(name).suffix.lower()
            if extension not in {".pdf", ".xlsx"}:
                continue
            content = archive.read(info)
            if extension == ".pdf":
                text, _, _, entries, _ = extract_pdf(content, name)
            else:
                text, _, _, entries, _ = extract_xlsx(content, name)
            texts.append(text)
            provenance.extend(entries)
    text = "\n".join(texts)
    fields, lines, field_provenance = parse_text(text, filename)
    return (
        text,
        fields,
        lines,
        provenance + field_provenance,
        "PROCESSED" if text.strip() else "MANUAL_REVIEW",
    )


def extract(data: bytes, extension: str, filename: str):
    if extension == ".pdf":
        return extract_pdf(data, filename)
    if extension == ".xlsx":
        return extract_xlsx(data, filename)
    return extract_zip(data, filename)


def infer_kind(filename: str) -> DocumentKind:
    stem = filename.lower()
    if "invoice" in stem:
        return "invoice"
    if "receipt" in stem:
        return "receipt"
    if "delivery" in stem:
        return "delivery_note"
    if "po" in stem or "purchase" in stem:
        return "po"
    return "unknown"
